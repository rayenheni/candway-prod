"""Regression tests: rubric draft endpoints must persist ``criteria_json`` as text.

``Rubric.criteria_json`` is a ``Text`` column. Several endpoints used to
assign a Python ``dict`` to it directly (``model_dump()`` / raw request
bodies). PyMySQL rejects that outright (``TypeError: dict can not be used as
parameter``) and SQLite raises ``ProgrammingError``, so the endpoints 500'd:

* ``POST /rubric/duplicate/{job_id}``  (called by the React rubric UI)
* ``POST /rubric/drafts/{job_id}``
* ``PUT  /rubric/drafts/{draft_id}``

It also covers the endpoints that read those rows back (draft read/list,
publish, Excel export of a draft-only job) and the default-rubric creation
path, which must never insert a rubric row without a company.
"""

import copy
import io
import json

import openpyxl
import pytest

from backend.database import Job, User
from backend.database import Rubric as RubricDB
from backend.dependencies import pwd_context
from backend.rubric.rubric_loader import _create_default_rubric, invalidate_cache
from backend.rubric.rubric_schema import JobRubric


@pytest.fixture
def published_job(db_session, test_company, test_recruiter):
    job = Job(
        title="Backend Engineer",
        recruiter_id=test_recruiter.id,
        company_id=test_company.id,
    )
    db_session.add(job)
    db_session.flush()
    rubric = JobRubric(job_id=job.id, version=1, categories=[])
    db_session.add(
        RubricDB(
            job_id=job.id,
            version=1,
            is_active=1,
            criteria_json=rubric.model_dump_json(),
            created_by=test_recruiter.id,
            company_id=test_company.id,
        )
    )
    db_session.commit()
    invalidate_cache(job.id)
    yield job
    invalidate_cache(job.id)


def _stored_criteria(db_session, rubric_id):
    db_session.expire_all()
    row = db_session.query(RubricDB).filter(RubricDB.id == rubric_id).one()
    assert isinstance(row.criteria_json, str), type(row.criteria_json)
    return json.loads(row.criteria_json)


def test_duplicate_rubric_persists_json_text(
    client, db_session, recruiter_headers, published_job
):
    resp = client.post(
        f"/api/v1/rubric/duplicate/{published_job.id}", headers=recruiter_headers
    )
    assert resp.status_code == 200, resp.text
    draft_id = resp.json()["id"]
    data = _stored_criteria(db_session, draft_id)
    assert data["job_id"] == published_job.id


def test_create_and_update_draft_persist_json_text(
    client, db_session, recruiter_headers, published_job
):
    resp = client.post(
        f"/api/v1/rubric/drafts/{published_job.id}",
        json={"name": "My draft"},
        headers=recruiter_headers,
    )
    assert resp.status_code == 200, resp.text
    draft_id = resp.json()["id"]
    assert _stored_criteria(db_session, draft_id)["job_id"] == published_job.id

    new_body = {"job_id": published_job.id, "version": 1, "categories": []}
    resp = client.put(
        f"/api/v1/rubric/drafts/{draft_id}",
        json={"name": "Renamed", "rubric_json": new_body},
        headers=recruiter_headers,
    )
    assert resp.status_code == 200, resp.text
    assert _stored_criteria(db_session, draft_id) == new_body


# ---------------------------------------------------------------------------
# Draft read / publish / export
# ---------------------------------------------------------------------------

VALID_CATEGORIES = [
    {
        "name": "Engineering",
        "weight": 1.0,
        "subcategories": [
            {
                "name": "Backend",
                "skills": [
                    {
                        "name": "API Design",
                        "level": "advanced",
                        "keywords": ["rest", "openapi"],
                    }
                ],
            }
        ],
    }
]


def _create_draft(client, headers, job_id, categories=None):
    resp = client.post(
        f"/api/v1/rubric/drafts/{job_id}", json={"name": "Draft"}, headers=headers
    )
    assert resp.status_code == 200, resp.text
    draft_id = resp.json()["id"]
    if categories is not None:
        resp = client.put(
            f"/api/v1/rubric/drafts/{draft_id}",
            json={
                "rubric_json": {
                    "job_id": job_id,
                    "version": 1,
                    "categories": categories,
                }
            },
            headers=headers,
        )
        assert resp.status_code == 200, resp.text
    return draft_id


def _rows(db_session, job_id):
    db_session.expire_all()
    return (
        db_session.query(RubricDB)
        .filter(RubricDB.job_id == job_id)
        .order_by(RubricDB.id)
        .all()
    )


def test_get_and_list_draft(client, recruiter_headers, published_job):
    draft_id = _create_draft(
        client, recruiter_headers, published_job.id, VALID_CATEGORIES
    )

    resp = client.get(f"/api/v1/rubric/drafts/{draft_id}", headers=recruiter_headers)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["id"] == draft_id
    assert body["job_id"] == published_job.id
    assert json.loads(body["criteria_json"])["categories"] == VALID_CATEGORIES

    resp = client.get(
        "/api/v1/rubric/drafts",
        params={"job_id": published_job.id},
        headers=recruiter_headers,
    )
    assert resp.status_code == 200, resp.text
    # Only the draft is listed, never the published (is_active=1) rubric.
    assert [d["id"] for d in resp.json()] == [draft_id]


def test_draft_read_denies_other_company_unknown_and_candidates(
    client, recruiter_headers, recruiter_headers_b, auth_headers, published_job
):
    draft_id = _create_draft(client, recruiter_headers, published_job.id)

    # Another company's recruiter: indistinguishable from a missing draft.
    resp = client.get(f"/api/v1/rubric/drafts/{draft_id}", headers=recruiter_headers_b)
    assert resp.status_code == 404, resp.text
    resp = client.get("/api/v1/rubric/drafts", headers=recruiter_headers_b)
    assert resp.status_code == 200 and resp.json() == []

    resp = client.get("/api/v1/rubric/drafts/999999", headers=recruiter_headers)
    assert resp.status_code == 404

    # Candidates are not recruiters.
    resp = client.get(f"/api/v1/rubric/drafts/{draft_id}", headers=auth_headers)
    assert resp.status_code == 403


def test_publish_draft_creates_next_active_version(
    client, db_session, recruiter_headers, published_job, test_company
):
    draft_id = _create_draft(
        client, recruiter_headers, published_job.id, VALID_CATEGORIES
    )

    resp = client.post(
        f"/api/v1/rubric/drafts/{draft_id}/publish", json={}, headers=recruiter_headers
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["version"] == 2
    assert body["draft_id"] == draft_id

    rows = {r.id: r for r in _rows(db_session, published_job.id)}
    active = [r for r in rows.values() if r.is_active == 1]
    assert [r.id for r in active] == [body["rubric_id"]]
    new = rows[body["rubric_id"]]
    assert new.version == 2
    assert new.company_id == rows[draft_id].company_id == test_company.id
    assert isinstance(new.criteria_json, str)
    assert json.loads(new.criteria_json)["categories"][0]["name"] == "Engineering"
    # The previous published version is kept but deactivated.
    v1 = [r for r in rows.values() if r.version == 1]
    assert len(v1) == 1 and v1[0].is_active == 0


def test_publish_draft_rejects_invalid_criteria_without_side_effects(
    client, db_session, recruiter_headers, published_job
):
    draft_id = _create_draft(
        client, recruiter_headers, published_job.id, [{"weight": "not-a-number"}]
    )
    before = [
        (r.id, r.version, r.is_active) for r in _rows(db_session, published_job.id)
    ]

    resp = client.post(
        f"/api/v1/rubric/drafts/{draft_id}/publish", json={}, headers=recruiter_headers
    )
    assert resp.status_code == 400, resp.text
    assert "Invalid rubric" in resp.json()["detail"]
    after = [
        (r.id, r.version, r.is_active) for r in _rows(db_session, published_job.id)
    ]
    assert after == before


def test_publish_draft_denies_other_company_and_unknown(
    client, db_session, recruiter_headers, recruiter_headers_b, published_job
):
    draft_id = _create_draft(
        client, recruiter_headers, published_job.id, VALID_CATEGORIES
    )
    before = [(r.id, r.is_active) for r in _rows(db_session, published_job.id)]

    resp = client.post(
        f"/api/v1/rubric/drafts/{draft_id}/publish",
        json={},
        headers=recruiter_headers_b,
    )
    assert resp.status_code == 404, resp.text
    assert [(r.id, r.is_active) for r in _rows(db_session, published_job.id)] == before

    resp = client.post(
        "/api/v1/rubric/drafts/999999/publish", json={}, headers=recruiter_headers
    )
    assert resp.status_code == 404


@pytest.fixture
def admin_headers(client, db_session):
    """/rubric/export requires the admin-only manage_content permission."""
    db_session.add(
        User(
            email="rubric-admin@example.com",
            name="Rubric Admin",
            hashed_password=pwd_context.hash("rubricadminpass123"),
            role="admin",
            is_super_admin=True,
            email_verified=True,
        )
    )
    db_session.commit()
    resp = client.get("/login")
    csrf = resp.headers.get("X-CSRF-Token") or resp.cookies.get("csrf_token") or ""
    resp = client.post(
        "/api/v1/auth/login",
        json={"email": "rubric-admin@example.com", "password": "rubricadminpass123"},
        headers={"X-CSRF-Token": csrf},
    )
    assert resp.status_code == 200, resp.text
    return {
        "Authorization": f"Bearer {resp.json()['access_token']}",
        "X-CSRF-Token": csrf,
    }


def test_export_decodes_draft_only_rubric(
    client, db_session, admin_headers, test_company, test_recruiter
):
    """A job with only a draft exports that draft; criteria_json is JSON text
    and used to be handed to the workbook builder undecoded (500)."""
    job = Job(
        title="Draft Only", recruiter_id=test_recruiter.id, company_id=test_company.id
    )
    db_session.add(job)
    db_session.flush()
    draft_json = {
        "job_id": job.id,
        "version": 1,
        "categories": copy.deepcopy(VALID_CATEGORIES),
    }
    db_session.add(
        RubricDB(
            job_id=job.id,
            version=0,
            is_active=0,
            criteria_json=json.dumps(draft_json),
            created_by=test_recruiter.id,
            company_id=test_company.id,
        )
    )
    db_session.commit()

    resp = client.get(f"/api/v1/rubric/export/{job.id}", headers=admin_headers)
    assert resp.status_code == 200, resp.text
    ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
    values = {c.value for row in ws.iter_rows() for c in row if c.value}
    assert "Engineering" in values
    assert any("API Design" in str(v) for v in values)


# ---------------------------------------------------------------------------
# Admin Excel import: the template must belong to an explicit company
# ---------------------------------------------------------------------------


def _template_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        ["Categorie", "Sous-categories", "Roles", "Criteres", "Competences", "Methodes"]
    )
    ws.append(
        ["Engineering", "Backend", "Developer", "Quality", "Python, SQL", "Live coding"]
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(client, headers, data):
    return client.post(
        "/api/v1/rubric/import",
        files={
            "file": (
                "Backend Template.xlsx",
                _template_xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data=data,
        headers=headers,
    )


def test_import_requires_company_id(client, db_session, admin_headers):
    before = db_session.query(RubricDB).count()
    resp = _import(client, admin_headers, {})
    assert resp.status_code == 400, resp.text
    assert "company_id" in resp.json()["detail"]
    assert db_session.query(RubricDB).count() == before


def test_import_rejects_unknown_or_inactive_company(
    client, db_session, admin_headers, test_company
):
    before = db_session.query(RubricDB).count()
    resp = _import(client, admin_headers, {"company_id": "987654"})
    assert resp.status_code == 404, resp.text

    test_company.is_active = False
    db_session.commit()
    resp = _import(client, admin_headers, {"company_id": str(test_company.id)})
    assert resp.status_code == 404, resp.text
    assert db_session.query(RubricDB).count() == before


def test_import_creates_company_owned_standalone_draft(
    client, db_session, admin_headers, test_company
):
    resp = _import(client, admin_headers, {"company_id": str(test_company.id)})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["categories_count"] == 1 and body["skills_count"] == 2

    db_session.expire_all()
    row = db_session.query(RubricDB).filter(RubricDB.id == body["draft_id"]).one()
    assert row.company_id == test_company.id
    # Standalone template: NULL, never 0 (job_id is a FK enforced by InnoDB).
    assert row.job_id is None
    assert row.is_active == 0 and row.version == 0
    assert isinstance(row.criteria_json, str)
    criteria = json.loads(row.criteria_json)
    skills = criteria["categories"][0]["subcategories"][0]["skills"]
    assert [s["name"] for s in skills] == ["Python", "SQL"]


# ---------------------------------------------------------------------------
# Default rubric creation (rubric_loader)
# ---------------------------------------------------------------------------


def test_default_rubric_is_owned_by_job_company(
    db_session, test_company, test_recruiter
):
    job = Job(
        title="No Rubric", recruiter_id=test_recruiter.id, company_id=test_company.id
    )
    db_session.add(job)
    db_session.flush()

    rubric = _create_default_rubric(job.id, db_session)

    assert rubric.job_id == job.id
    rows = _rows(db_session, job.id)
    assert len(rows) == 1
    assert rows[0].company_id == test_company.id
    assert rows[0].is_active == 1 and rows[0].version == 1


def test_default_rubric_without_resolvable_company_is_not_persisted(db_session):
    # No Job row -> no owning company. Must not attempt a NOT NULL-violating
    # insert (which used to fail the whole request with IntegrityError).
    rubric = _create_default_rubric(424242, db_session)

    assert rubric.job_id == 424242
    assert rubric.categories == []
    assert _rows(db_session, 424242) == []
    db_session.commit()  # session remains usable
